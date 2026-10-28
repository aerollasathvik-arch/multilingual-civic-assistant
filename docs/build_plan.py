#!/usr/bin/env python3
"""Build the 12-week project plan workbook (docs/project-plan.xlsx).

The workbook is the source of truth for task status. Edit TASKS here (or the
Status column in Excel) as weeks progress.
"""
import datetime as dt
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = pathlib.Path(__file__).with_name("project-plan.xlsx")
START = dt.date(2026, 10, 5)  # Monday of week 1 (week 4 = 26 Oct - 1 Nov 2026)
CURRENT_WEEK = 4

FONT = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUB = Font(name=FONT, size=10, italic=True, color="595959")
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
BAND = PatternFill("solid", fgColor="F2F5FA")
NOW = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {"Done": "C6EFCE", "In progress": "FFEB9C", "Not started": "FFFFFF", "Blocked": "FFC7CE"}

WEEKS = [
    (1, "Scope and research design",
     "Research question, objectives, scope (England), target languages, metrics, ethics pre-screen, repo"),
    (2, "Literature and architecture",
     "Verified reading list, system architecture, tool and licence choices"),
    (3, "Corpus collection",
     "38 official pages ingested with raw snapshots and a data card"),
    (4, "Chunking, index and retrieval baseline",
     "Section chunks, multilingual FAISS index, 150-query retrieval experiment, failure analysis"),
    (5, "Translation layer and retrieval hardening",
     "Jurisdiction filter, language ID, translate-then-retrieve vs direct comparison, verified translations"),
    (6, "Grounded answer generation",
     "LLM answers only from retrieved sources with numbered citations, plain-language style, refusal rule"),
    (7, "Baseline and evaluation set",
     "Ungrounded LLM baseline, ~100-question gold set incl. unanswerable, scoring rubric, ethics approval"),
    (8, "Accessible front-end",
     "Web prototype with language picker, right-to-left support, citation display; participants recruited"),
    (9, "Evaluation runs and user study",
     "Grounded vs ungrounded scored; small user study run; corpus re-crawled to measure content drift"),
    (10, "Analysis and results write-up",
     "Results and discussion chapters with tables, figures and failure cases"),
    (11, "Full draft",
     "Introduction, literature review, method, ethics and limitations chapters; full draft to supervisor"),
    (12, "Revision and submission",
     "Supervisor feedback applied, demo video, packaged artefact, final submission"),
]

# (week, id, task, output / evidence, status)
TASKS = [
    (1, "T01", "Fix research question, aim, objectives and scope (England only; health, housing, schools)", "docs/project-scope.md", "Done"),
    (1, "T02", "Choose target languages from evidence (ONS Census 2021 main languages)", "docs/project-scope.md", "Done"),
    (1, "T03", "Define metrics: retrieval hit@k/MRR, answer accuracy, hallucination rate, citation correctness, comprehension", "docs/evaluation-plan.md", "Done"),
    (1, "T04", "Ethics pre-screen for the user study and draft application", "docs/ethics-notes.md", "In progress"),
    (1, "T05", "Project repo, Python environment, pinned requirements", "requirements.txt", "Done"),
    (2, "T06", "Reading list on RAG, multilingual embeddings, MT and RAG evaluation (all references verified on arXiv)", "docs/literature-notes.md", "Done"),
    (2, "T07", "Architecture: ingest > chunk > embed > index > retrieve > generate > translate > UI", "docs/architecture.md", "Done"),
    (2, "T08", "Tool selection with licences (e5 MIT, MiniLM Apache-2.0, FAISS MIT, NLLB CC-BY-NC-4.0)", "docs/architecture.md", "Done"),
    (2, "T09", "Source licensing check (Open Government Licence v3.0; nhs.uk reuse terms)", "docs/data-card.md", "Done"),
    (3, "T10", "Curate and live-validate source list: 26 GOV.UK + 12 NHS.uk pages", "src/sources.py", "Done"),
    (3, "T11", "Ingestion: GOV.UK Content API + NHS HTML, raw snapshots, fetch timestamp", "src/ingest.py", "Done"),
    (3, "T12", "Handle GOV.UK formats (guide, answer, local_transaction, transaction, detailed_guide) and NHS lead text", "src/ingest.py", "Done"),
    (3, "T13", "Data card: counts per domain, dates, currency caveats (Renters' Rights Act, 1 May 2026)", "docs/data-card.md", "Done"),
    (4, "T14", "Heading-aware chunking (40-200 words, 40-word overlap, title + heading prefix)", "src/chunk.py", "Done"),
    (4, "T15", "Multilingual embeddings (multilingual-e5-base) + FAISS index", "src/retrieval.py", "Done"),
    (4, "T16", "Test set: 30 questions x 5 languages, verbatim answer phrase checked on gold page", "eval/queries.csv", "Done"),
    (4, "T17", "Retrieval experiment: BM25 vs notebook set-up vs e5 sections vs MiniLM sections", "results/retrieval_by_lang.csv", "Done"),
    (4, "T18", "Failure analysis (wrong nation, strict answer phrase, Urdu weakest)", "docs/week-04-report.md", "Done"),
    (4, "T19", "Rebuild week 4 notebook on the project pipeline", "notebooks/week04_retrieval.ipynb", "Done"),
    (4, "T20", "Weekly progress report for supervisor", "docs/week-04-report.md", "Done"),
    (5, "T21", "Tag chunks by UK nation and filter to England by default", "", "Not started"),
    (5, "T22", "Native-speaker check of pl/ro/ar/ur test questions", "", "Not started"),
    (5, "T23", "Language identification of the user's question", "", "Not started"),
    (5, "T24", "NLLB query translation; compare translate-then-retrieve vs direct multilingual retrieval", "", "Not started"),
    (5, "T25", "Try a reranker on top-20 to close the page-vs-answer gap (hit@3 0.96 vs ans@3 0.67)", "", "Not started"),
    (6, "T26", "Choose LLM API (multilingual quality, cost, data handling)", "", "Not started"),
    (6, "T27", "Grounded prompt: answer only from sources, numbered citations with URLs, plain language", "", "Not started"),
    (6, "T28", "Refusal rule for low retrieval score and out-of-scope questions", "", "Not started"),
    (6, "T29", "Answer in user's language: LLM direct vs NLLB back-translation", "", "Not started"),
    (7, "T30", "Ungrounded baseline: same LLM, no retrieved sources", "", "Not started"),
    (7, "T31", "Gold set to ~100 questions with reference answers, incl. unanswerable ones", "", "Not started"),
    (7, "T32", "Scoring rubric for accuracy, hallucination, citation correctness; second rater on a sample", "", "Not started"),
    (7, "T33", "Ethics approval confirmed before any participant contact", "", "Not started"),
    (8, "T34", "Web front-end: language picker, right-to-left layout, citation cards, last-updated dates", "", "Not started"),
    (8, "T35", "Accessibility pass (WCAG 2.2 AA basics, plain-language UI copy)", "", "Not started"),
    (8, "T36", "Recruit user study participants", "", "Not started"),
    (9, "T37", "Run grounded vs ungrounded on full gold set and score", "", "Not started"),
    (9, "T38", "User study sessions and comprehension questionnaire", "", "Not started"),
    (9, "T39", "Re-crawl corpus and measure content drift since week 3", "", "Not started"),
    (10, "T40", "Statistics, tables, figures; failure-case catalogue", "", "Not started"),
    (10, "T41", "Draft results and discussion chapters", "", "Not started"),
    (11, "T42", "Draft introduction, literature review, method, ethics and limitations", "", "Not started"),
    (11, "T43", "Full draft to supervisor", "", "Not started"),
    (12, "T44", "Apply feedback, record demo video, package artefact, submit", "", "Not started"),
]

RISKS = [
    ("Answer is correct for the wrong UK nation (Scotland/Wales/NI sections)", "High", "Seen in week 4 (school leaving age)", "Tag chunks by nation, filter to England, show jurisdiction in UI (T21)"),
    ("Official page changes after crawl (e.g. Renters' Rights Act from 1 May 2026)", "High", "public_updated_at did not reflect the 2026 changes", "Store fetch date, show it with every citation, re-crawl and measure drift (T39)"),
    ("Author-written translations bias the multilingual test set", "Medium", "All non-English queries unverified", "Native-speaker check before reporting final numbers (T22)"),
    ("Right page retrieved but not the answering section", "Medium", "hit@3 0.96 vs ans@3 0.67", "Reranker, better chunking, graded relevance judgments (T25, T32)"),
    ("Ethics approval delays the user study", "Medium", "Study planned for week 9", "Submit early; fall back to expert review if not approved"),
    ("NLLB licence is non-commercial", "Low", "CC-BY-NC-4.0", "Acceptable for research; note in deployment recommendations"),
    ("Council-specific rules (admissions dates, housing) live on 300+ council sites", "Medium", "GOV.UK local_transaction pages only redirect", "Assistant points users to their council; optional single-council case study"),
]


def header(ws, row, cols):
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill, c.border = H_FONT, H_FILL, BOX
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 24


def body(ws, r, ncols, fill=None):
    for i in range(1, ncols + 1):
        c = ws.cell(row=r, column=i)
        c.font = c.font if c.font and c.font.bold else BODY
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BOX
        if fill:
            c.fill = fill


def week_dates(w):
    s = START + dt.timedelta(weeks=w - 1)
    return s, s + dt.timedelta(days=6)


def main():
    wb = Workbook()

    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "A Multilingual AI Assistant for Accessing Health, Housing, and Schooling Information"
    ws["A1"].font = TITLE
    ws["A2"] = (f"12-week plan. Week 1 starts Mon {START:%d %b %Y}; current week is {CURRENT_WEEK} "
                "(highlighted). Task status lives on the Tasks sheet.")
    ws["A2"].font = SUB
    cols = [("Week", 7), ("Start", 12), ("End", 12), ("Theme", 34), ("Deliverable", 70), ("Done", 9), ("Tasks", 8)]
    header(ws, 4, cols)
    for i, (w, theme, deliv) in enumerate(WEEKS):
        r = 5 + i
        s, e = week_dates(w)
        ws.append([w, s, e, theme, deliv, None, None])
        ws.cell(row=r, column=6, value=f'=COUNTIFS(Tasks!$A:$A,A{r},Tasks!$E:$E,"Done")')
        ws.cell(row=r, column=7, value=f"=COUNTIF(Tasks!$A:$A,A{r})")
        for col in (2, 3):
            ws.cell(row=r, column=col).number_format = "ddd dd mmm yyyy"
        fill = NOW if w == CURRENT_WEEK else (BAND if i % 2 else None)
        body(ws, r, len(cols), fill)
    last = 5 + len(WEEKS)
    ws.cell(row=last, column=5, value="Total").font = BOLD
    ws.cell(row=last, column=6, value=f"=SUM(F5:F{last - 1})").font = BOLD
    ws.cell(row=last, column=7, value=f"=SUM(G5:G{last - 1})").font = BOLD
    ws.freeze_panes = "A5"

    ws = wb.create_sheet("Tasks")
    cols = [("Week", 7), ("ID", 7), ("Task", 78), ("Output / evidence", 34), ("Status", 13)]
    header(ws, 1, cols)
    dv = DataValidation(type="list", formula1='"Not started,In progress,Blocked,Done"', allow_blank=False)
    ws.add_data_validation(dv)
    for i, t in enumerate(TASKS):
        r = 2 + i
        ws.append(list(t))
        body(ws, r, len(cols), BAND if t[0] % 2 == 0 else None)
        st = ws.cell(row=r, column=5)
        st.fill = PatternFill("solid", fgColor=STATUS_FILL[t[4]])
        dv.add(st)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{1 + len(TASKS)}"

    ws = wb.create_sheet("Week 4 results")
    ws["A1"] = "Retrieval experiment, 150 queries (30 questions x en/pl/ro/ar/ur)"
    ws["A1"].font = TITLE
    ws["A2"] = "Source: results/retrieval_by_lang.csv. ans@3 = a top-3 chunk contains the verbatim answer phrase (strict lower bound)."
    ws["A2"].font = SUB
    import pandas as pd  # local import: only this sheet needs it
    res = pd.read_csv(pathlib.Path(__file__).parents[1] / "results/retrieval_by_lang.csv")
    cols = [("System", 20), ("Language", 10), ("hit@1", 9), ("hit@3", 9), ("MRR@10", 9), ("ans@3", 9)]
    header(ws, 4, cols)
    for i, row in enumerate(res.itertuples(index=False)):
        ws.append(list(row))
        body(ws, 5 + i, len(cols), NOW if row.lang == "ALL" else None)
        for c in range(3, 7):
            ws.cell(row=5 + i, column=c).number_format = "0.000"

    ws = wb.create_sheet("Risks")
    cols = [("Risk", 52), ("Severity", 10), ("Evidence", 40), ("Mitigation", 60)]
    header(ws, 1, cols)
    for i, r in enumerate(RISKS):
        ws.append(list(r))
        body(ws, 2 + i, len(cols))

    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
